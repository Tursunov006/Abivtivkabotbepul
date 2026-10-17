import os
from docx import Document
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from texts import DOC


def set_cell_margins(cell, top=40, bottom=40, left=40, right=40):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for edge, val in borders.items():
        tag = f'w:{edge}'
        el = OxmlElement(tag)
        if val is not None:
            el.set(qn('w:val'), val.get('val', 'single'))
            el.set(qn('w:sz'), str(val.get('sz', 4)))
            el.set(qn('w:space'), '0')
            el.set(qn('w:color'), val.get('color', '000000'))
        else:
            el.set(qn('w:val'), 'none')
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_table_borders(table, color="000000"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)


def generate_docx(data: dict, output_path: str):
    lang = data.get("lang", "lat")
    if lang not in DOC:
        lang = "lat"
    labels = DOC[lang]["labels"]
    none_text = DOC[lang]["none"]

    doc = Document()

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(20)
    section.right_margin = Mm(15)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # 1. Yuqori blok: Chapda (MA'LUMOTNOMA, F.I.Sh., Universitet/fakultet), O'ngda 3x4 rasm
    top_table = doc.add_table(rows=1, cols=2)
    top_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    top_table.autofit = False

    cell_info = top_table.cell(0, 0)
    cell_photo = top_table.cell(0, 1)

    cell_info.width = Mm(138)
    cell_photo.width = Mm(37)

    set_cell_borders(cell_info)
    set_cell_borders(cell_photo)
    set_cell_margins(cell_info, top=0, bottom=0, left=0, right=60)
    set_cell_margins(cell_photo, top=0, bottom=0, left=0, right=0)

    # MA'LUMOTNOMA
    p_title = cell_info.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run(DOC[lang]["title"])
    r_title.bold = True
    r_title.font.size = Pt(12.5)

    # F.I.Sh.
    p_fio = cell_info.add_paragraph()
    p_fio.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fio.paragraph_format.space_before = Pt(0)
    p_fio.paragraph_format.space_after = Pt(8)
    r_fio = p_fio.add_run(data.get("fio", "").strip())
    r_fio.bold = True
    r_fio.font.size = Pt(12.5)

    # Universitet, fakultet, kurs
    pos_val = data.get("position", "").strip()
    if pos_val:
        p_pos = cell_info.add_paragraph()
        p_pos.paragraph_format.space_before = Pt(0)
        p_pos.paragraph_format.space_after = Pt(6)
        r_pos = p_pos.add_run(pos_val)
        r_pos.font.size = Pt(10.5)

    # 3x4 Rasm
    photo_path = data.get("photo_path")
    p_photo = cell_photo.paragraphs[0]
    p_photo.paragraph_format.space_before = Pt(0)
    p_photo.paragraph_format.space_after = Pt(0)

    if photo_path and os.path.exists(photo_path):
        try:
            p_photo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_photo.add_run().add_picture(photo_path, width=Mm(33), height=Mm(43))
        except Exception:
            p_photo.add_run("3x4")
    else:
        p_photo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        box_border = {"sz": 6, "val": "single", "color": "4A6984"}
        set_cell_borders(cell_photo, top=box_border, bottom=box_border, left=box_border, right=box_border)
        set_cell_margins(cell_photo, top=280, bottom=280, left=40, right=40)
        p_photo.add_run("\n3x4\n")

    # 2. Asosiy bandlar (Tug'ilgan yili, joyi, Millati, Ta'lim yo'nalishi, Tamomlagan, Telefon, Chet tillari)
    b_date = data.get("birth_date", "").strip() or none_text
    b_place = data.get("birth_place", "").strip() or none_text
    nat_val = data.get("nationality", "").strip() or none_text
    spec_val = data.get("specialty", "").strip() or none_text
    grad_val = data.get("graduated", "").strip() or none_text
    phone_val = data.get("contact", "").strip() or none_text
    lang_val = data.get("languages", "").strip() or none_text
    awards_val = data.get("awards", "").strip() or none_text

    fields_spec = [
        (labels["birth_date"], b_date),
        (labels["birth_place"], b_place),
        (labels["nationality"], nat_val),
        (labels["specialty"], spec_val),
        (labels["graduated"], grad_val),
        (labels["contact"], phone_val),
        (labels["languages"], lang_val),
    ]

    info_table = doc.add_table(rows=0, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    for idx, (lbl, val) in enumerate(fields_spec):
        row = info_table.add_row()
        c_lbl = row.cells[0]
        c_val = row.cells[1]

        c_lbl.width = Mm(68)
        c_val.width = Mm(107)

        bottom_pad = 90 if idx == 4 else 35  # Tamomlagan dan keyin biroz bo'shliq
        set_cell_margins(c_lbl, top=30, bottom=bottom_pad, left=0, right=40)
        set_cell_margins(c_val, top=30, bottom=bottom_pad, left=0, right=0)
        set_cell_borders(c_lbl)
        set_cell_borders(c_val)

        p1 = c_lbl.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(lbl)
        r1.bold = True
        r1.font.size = Pt(10.2)

        p2 = c_val.paragraphs[0]
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(1)
        r2 = p2.add_run(val)
        r2.font.size = Pt(10.2)

    # Davlat mukofotlari qatori
    p_aw = doc.add_paragraph()
    p_aw.paragraph_format.space_before = Pt(3)
    p_aw.paragraph_format.space_after = Pt(8)
    r_aw_lbl = p_aw.add_run(f"{labels['awards']}          ")
    r_aw_lbl.bold = True
    r_aw_lbl.font.size = Pt(10.2)
    r_aw_val = p_aw.add_run(awards_val)
    r_aw_val.font.size = Pt(10.2)

    # 3. SHAXSIY MA'LUMOTLARI (Oila a'zolari jadvali - sarlavha qatorisiz)
    relatives = data.get("relatives", [])
    valid_relatives = [
        r for r in relatives
        if r.get("fio") or r.get("job") or r.get("birth") or r.get("address")
    ]

    if valid_relatives:
        p_rt = doc.add_paragraph()
        p_rt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_rt.paragraph_format.space_before = Pt(6)
        p_rt.paragraph_format.space_after = Pt(8)
        r_rt = p_rt.add_run(DOC[lang]["rel_title"])
        r_rt.bold = True
        r_rt.font.size = Pt(11)

        rtable = doc.add_table(rows=0, cols=5)
        rtable.alignment = WD_TABLE_ALIGNMENT.CENTER
        rtable.autofit = False
        set_table_borders(rtable, color="000000")

        col_widths = [Mm(22), Mm(41), Mm(26), Mm(50), Mm(36)]

        for rel in valid_relatives:
            row_cells = rtable.add_row().cells
            row_vals = [
                (rel.get("kind", "") or "").strip().upper(),
                (rel.get("fio", "") or "").strip(),
                (rel.get("birth", "") or "").strip(),
                (rel.get("address", "") or "").strip(),
                (rel.get("job", "") or "").strip()
            ]
            for i, val in enumerate(row_vals):
                row_cells[i].width = col_widths[i]
                p = row_cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(val)
                r.font.size = Pt(9.8)
                if i == 0:
                    r.bold = True
                set_cell_margins(row_cells[i], top=45, bottom=55, left=55, right=55)

    # 4. ERISHGAN YUTUQLARI
    ach_raw = (data.get("achievements", "") or data.get("work", "")).strip()
    if ach_raw and ach_raw.lower() not in ["yo'q", "yo‘q", "йўқ", "нет", "-"]:
        p_at = doc.add_paragraph()
        p_at.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_at.paragraph_format.space_before = Pt(12)
        p_at.paragraph_format.space_after = Pt(8)
        r_at = p_at.add_run(DOC[lang]["ach_title"])
        r_at.bold = True
        r_at.font.size = Pt(11)

        ach_lines = [l.strip() for l in ach_raw.split("\n") if l.strip()]
        for idx, line in enumerate(ach_lines, start=1):
            clean_line = line
            if len(clean_line) > 2 and clean_line[0].isdigit() and clean_line[1] in ['.', ')']:
                clean_line = clean_line[2:].strip()
            elif len(clean_line) > 3 and clean_line[:2].isdigit() and clean_line[2] in ['.', ')']:
                clean_line = clean_line[3:].strip()

            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Mm(8)
            p_item.paragraph_format.first_line_indent = Mm(-5)
            p_item.paragraph_format.space_before = Pt(2)
            p_item.paragraph_format.space_after = Pt(6)
            r_item = p_item.add_run(f"{idx}.  {clean_line}")
            r_item.font.size = Pt(10)

    doc.save(output_path)
    return output_path
